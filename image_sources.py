import re
from pathlib import Path

try:
    import pymupdf
except Exception:
    pymupdf = None

BASE = Path(__file__).resolve().parent
EXCEL_IMG_SRC = "sources/NAUTICA ATS TOM .xlsx"
EXCEL_IMG_DIR = "imagenes_excel"
PDF_IMG_SRC = "sources/MMX FALL ATS 6.23 LEVY GROUP 3.pdf"
PDF_IMG_SRC2 = "sources/LAUNDRY EARLY BUY LINESHEET 3.18.pdf"
PDF_IMG_DIR = "imagenes_pdf"
XLS_IMG_DIR = "imagenes_xls"
ATS_IMG_DIR = "imagenes_ats"

MANUAL_IMAGE_OVERRIDES = {
    ("V322453M", "STO"): "V322453M_COLOR_STONE.png",
}

STYLE_IMAGE_OVERRIDES = {
    "119004414": "imagenes/119004414.png",
    "117004093": "imagenes/117004093.jpg",
    "119000214": "imagenes/119000214.jpg",
    "119000455": "imagenes/119000455.jpg",
    "119001566": "imagenes/119001566.jpg",
    "119001328": "imagenes/119001328.jpg",
    "119005385": "imagenes/119005385.jpg",
    "121000428": "imagenes/121000428.jpg",
}


def safe_name(s):
    s = re.sub(r'[\\/*?:"<>|]', "_", str(s))
    s = re.sub(r"\s+", "_", s).strip("_")
    return (s[:50] or "item")


def extract_excel_images():
    src = BASE / EXCEL_IMG_SRC
    out_dir = BASE / EXCEL_IMG_DIR
    out_dir.mkdir(exist_ok=True)
    mapping = {}
    if not src.exists():
        return mapping
    from openpyxl import load_workbook
    wb = load_workbook(src)
    ws = wb["Sheet"] if "Sheet" in wb.sheetnames else wb.worksheets[1]
    for im in ws._images:
        anchor = im.anchor
        from_cell = getattr(anchor, "_from", None)
        if from_cell is None:
            continue
        excel_row = from_cell.row + 1
        tl = ws.cell(row=excel_row, column=1).value
        style = ws.cell(row=excel_row, column=2).value
        if tl is None or style is None:
            continue
        style_s = str(style).strip()
        if "total" in style_s.lower():
            continue
        key = (str(tl).strip(), style_s)
        fname = f"excel_{safe_name(tl)}_{safe_name(style)}.png"
        dest = out_dir / fname
        if not dest.exists():
            try:
                dest.write_bytes(im._data())
            except Exception:
                continue
        mapping[key] = fname
    return mapping


def extract_pdf_images():
    src = BASE / PDF_IMG_SRC
    out_dir = BASE / PDF_IMG_DIR
    out_dir.mkdir(exist_ok=True)
    mapping = {}
    if pymupdf is None or not src.exists():
        return mapping
    doc = pymupdf.open(src)
    candidates = {}
    for pno in range(doc.page_count):
        page = doc[pno]
        style_blocks = []
        for b in page.get_text("blocks"):
            x0, y0, x1, y1, txt = b[:5]
            txt = txt.strip()
            m = re.search(r"\b(\d{7,})\b", txt)
            if m and x0 > 180:
                style_blocks.append((y0, y1, m.group(1)))
        for img in page.get_images(full=True):
            for r in page.get_image_rects(img[0]):
                styles = {sb[2] for sb in style_blocks if not (sb[1] < r.y0 or sb[0] > r.y1)}
                for st in styles:
                    candidates.setdefault(st, []).append((r.x0, r.y0, img[0]))
    for style, cands in candidates.items():
        cands_sorted = sorted(cands, key=lambda c: (c[0], c[1]))
        _, _, xref = cands_sorted[0]
        base = doc.extract_image(xref)
        ext = base.get("ext", "png")
        fname = f"pdf_{safe_name(style)}.{ext}"
        dest = out_dir / fname
        if not dest.exists():
            try:
                dest.write_bytes(base["image"])
            except Exception:
                continue
        mapping[style] = fname
    doc.close()
    return mapping


def extract_pdf_images2():
    src = BASE / PDF_IMG_SRC2
    out_dir = BASE / PDF_IMG_DIR
    out_dir.mkdir(exist_ok=True)
    mapping = {}
    if pymupdf is None or not src.exists():
        return mapping
    doc = pymupdf.open(src)
    used = set()
    for pno in range(doc.page_count):
        page = doc[pno]
        styles = []
        for b in page.get_text("blocks"):
            x0, y0, x1, y1, txt = b[:5]
            m = re.search(r"\b(\d{7,})\b", txt.strip())
            if m:
                styles.append((m.group(1), (x0 + x1) / 2, y0, y1))
        imgs = []
        for img in page.get_images(full=True):
            for r in page.get_image_rects(img[0]):
                imgs.append((img[0], (r.x0 + r.x1) / 2, r.y0, r.y1))
        for style, scx, sy0, sy1 in styles:
            cands = []
            for xref, icx, iy0, iy1 in imgs:
                if xref in used:
                    continue
                if iy0 >= sy1 - 5 and abs(icx - scx) <= 90:
                    cands.append((abs(iy0 - sy1) + abs(icx - scx) * 0.5, xref))
            if cands:
                cands.sort()
                xref = cands[0][1]
                used.add(xref)
                base = doc.extract_image(xref)
                ext = base.get("ext", "png")
                fname = f"pdf_{safe_name(style)}.{ext}"
                dest = out_dir / fname
                if not dest.exists():
                    try:
                        dest.write_bytes(base["image"])
                    except Exception:
                        continue
                mapping[style] = fname
    doc.close()
    return mapping


def scan_pdf_folder():
    mapping = {}
    out_dir = BASE / PDF_IMG_DIR
    if not out_dir.exists():
        return mapping
    for f in out_dir.glob("pdf_*"):
        if f.suffix.lower() in (".png", ".jpg", ".jpeg"):
            mapping[f.stem[4:]] = f.name
    return mapping


def extract_xls_images():
    out_dir = BASE / XLS_IMG_DIR
    mapping = {}
    if not out_dir.exists():
        return mapping
    for f in out_dir.iterdir():
        if not f.is_file():
            continue
        name = f.name
        if not name.startswith("xls_") or not name.lower().endswith((".png", ".jpg", ".jpeg")):
            continue
        key = name[4:].rsplit(".", 1)[0]
        mapping[key] = f"{XLS_IMG_DIR}/{name}"
    if "V623040" in mapping:
        mapping.setdefault("V623040A", mapping["V623040"])
        mapping.setdefault("V623040H", mapping["V623040"])
    return mapping


def extract_ats_images():
    out_dir = BASE / ATS_IMG_DIR
    mapping = {}
    if not out_dir.exists():
        return mapping
    by_style = {}
    for f in out_dir.iterdir():
        if not f.is_file():
            continue
        name = f.name
        if not name.startswith("ATS_") or not name.lower().endswith((".png", ".jpg", ".jpeg")):
            continue
        stem = name[4:].rsplit(".", 1)[0]
        parts = stem.split("_")
        if len(parts) < 2:
            continue
        style = parts[0]
        idx = 0
        if len(parts) >= 3 and parts[-1].isdigit():
            idx = int(parts[-1])
        by_style.setdefault(style, []).append((idx, name))
    for style, cands in by_style.items():
        cands.sort()
        mapping[style] = f"{ATS_IMG_DIR}/{cands[0][1]}"
    return mapping


def extract_imagenes():
    out_dir = BASE / "imagenes"
    mapping = {}
    if not out_dir.exists():
        return mapping
    for f in out_dir.iterdir():
        if not f.is_file():
            continue
        if f.suffix.lower() not in (".png", ".jpg", ".jpeg"):
            continue
        stem = f.stem
        if stem not in mapping:
            mapping[stem] = f"imagenes/{f.name}"
    return mapping


def build_maps():
    pdf_map = extract_pdf_images()
    for k, v in extract_pdf_images2().items():
        pdf_map.setdefault(k, v)
    for k, v in scan_pdf_folder().items():
        pdf_map.setdefault(k, v)
    return {
        "excel": extract_excel_images(),
        "pdf": pdf_map,
        "xls": extract_xls_images(),
        "ats": extract_ats_images(),
        "imagenes": extract_imagenes(),
    }


def combined_image(r, maps):
    _style = str(r.get("Style", "")).strip()
    _color = str(r.get("Color_code", "")).strip().upper()
    if (_style, _color) in MANUAL_IMAGE_OVERRIDES:
        return MANUAL_IMAGE_OVERRIDES[(_style, _color)]
    if _style in STYLE_IMAGE_OVERRIDES:
        return STYLE_IMAGE_OVERRIDES[_style]
    excel_fname = maps["excel"].get(
        (str(r.get("BRAND_NAME", "")).strip(), _style), ""
    )
    if excel_fname:
        return f"{EXCEL_IMG_DIR}/{excel_fname}"
    pdf_fname = maps["pdf"].get(_style, "")
    if pdf_fname:
        return f"{PDF_IMG_DIR}/{pdf_fname}"
    xls_fname = maps["xls"].get(_style, "")
    if xls_fname:
        return xls_fname
    ats_fname = maps["ats"].get(_style, "")
    if ats_fname:
        return ats_fname
    img_fname = maps["imagenes"].get(_style, "")
    if img_fname:
        return img_fname
    return ""
