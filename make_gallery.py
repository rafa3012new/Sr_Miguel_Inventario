import html
import re
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

try:
    import pymupdf
except Exception:
    pymupdf = None

BASE = Path(__file__).resolve().parent
IMAGE_COL = "Imagen"
EXCEL_IMG_SRC = "NAUTICA ATS TOM .xlsx"
EXCEL_IMG_DIR = "imagenes_excel"
PDF_IMG_SRC = "MMX FALL ATS 6.23 LEVY GROUP 3.pdf"
PDF_IMG_SRC2 = "LAUNDRY EARLY BUY LINESHEET 3.18.pdf"
PDF_IMG_DIR = "imagenes_pdf"
XLS_IMG_DIR = "imagenes_xls"
ATS_IMG_DIR = "imagenes_ats"


def safe_name(s):
    s = re.sub(r'[\\/*?:"<>|]', "_", str(s))
    s = re.sub(r"\s+", "_", s).strip("_")
    return (s[:50] or "item")


def extract_excel_images():
    src = BASE / EXCEL_IMG_SRC
    out_dir = BASE / EXCEL_IMG_DIR
    out_dir.mkdir(exist_ok=True)
    mapping = {}
    if not src.exists():
        return mapping
    wb = load_workbook(src)
    ws = wb["Sheet"] if "Sheet" in wb.sheetnames else wb.worksheets[1]
    for im in ws._images:
        anchor = im.anchor
        from_cell = getattr(anchor, "_from", None)
        if from_cell is None:
            continue
        excel_row = from_cell.row + 1
        tl = ws.cell(row=excel_row, column=1).value
        style = ws.cell(row=excel_row, column=2).value
        if tl is None or style is None:
            continue
        style_s = str(style).strip()
        if "total" in style_s.lower():
            continue
        key = (str(tl).strip(), style_s)
        fname = f"excel_{safe_name(tl)}_{safe_name(style)}.png"
        dest = out_dir / fname
        if not dest.exists():
            try:
                dest.write_bytes(im._data())
            except Exception:
                continue
        mapping[key] = fname
    return mapping


def extract_pdf_images():
    src = BASE / PDF_IMG_SRC
    out_dir = BASE / PDF_IMG_DIR
    out_dir.mkdir(exist_ok=True)
    mapping = {}
    if pymupdf is None or not src.exists():
        return mapping
    doc = pymupdf.open(src)
    candidates = {}
    for pno in range(doc.page_count):
        page = doc[pno]
        style_blocks = []
        for b in page.get_text("blocks"):
            x0, y0, x1, y1, txt = b[:5]
            txt = txt.strip()
            m = re.search(r"\b(\d{7,})\b", txt)
            if m and x0 > 180:
                style_blocks.append((y0, y1, m.group(1)))
        for img in page.get_images(full=True):
            for r in page.get_image_rects(img[0]):
                styles = {sb[2] for sb in style_blocks if not (sb[1] < r.y0 or sb[0] > r.y1)}
                for st in styles:
                    candidates.setdefault(st, []).append((r.x0, r.y0, img[0]))
    for style, cands in candidates.items():
        cands_sorted = sorted(cands, key=lambda c: (c[0], c[1]))
        _, _, xref = cands_sorted[0]
        base = doc.extract_image(xref)
        ext = base.get("ext", "png")
        fname = f"pdf_{safe_name(style)}.{ext}"
        dest = out_dir / fname
        if not dest.exists():
            try:
                dest.write_bytes(base["image"])
            except Exception:
                continue
        mapping[style] = fname
    doc.close()
    return mapping


def extract_pdf_images2():
    src = BASE / PDF_IMG_SRC2
    out_dir = BASE / PDF_IMG_DIR
    out_dir.mkdir(exist_ok=True)
    mapping = {}
    if pymupdf is None or not src.exists():
        return mapping
    doc = pymupdf.open(src)
    used = set()
    for pno in range(doc.page_count):
        page = doc[pno]
        styles = []
        for b in page.get_text("blocks"):
            x0, y0, x1, y1, txt = b[:5]
            m = re.search(r"\b(\d{7,})\b", txt.strip())
            if m:
                styles.append((m.group(1), (x0 + x1) / 2, y0, y1))
        imgs = []
        for img in page.get_images(full=True):
            for r in page.get_image_rects(img[0]):
                imgs.append((img[0], (r.x0 + r.x1) / 2, r.y0, r.y1))
        for style, scx, sy0, sy1 in styles:
            cands = []
            for xref, icx, iy0, iy1 in imgs:
                if xref in used:
                    continue
                if iy0 >= sy1 - 5 and abs(icx - scx) <= 90:
                    cands.append((abs(iy0 - sy1) + abs(icx - scx) * 0.5, xref))
            if cands:
                cands.sort()
                xref = cands[0][1]
                used.add(xref)
                base = doc.extract_image(xref)
                ext = base.get("ext", "png")
                fname = f"pdf_{safe_name(style)}.{ext}"
                dest = out_dir / fname
                if not dest.exists():
                    try:
                        dest.write_bytes(base["image"])
                    except Exception:
                        continue
                mapping[style] = fname
    doc.close()
    return mapping


def extract_xls_images():
    out_dir = BASE / XLS_IMG_DIR
    mapping = {}
    if not out_dir.exists():
        return mapping
    for f in out_dir.iterdir():
        if not f.is_file():
            continue
        name = f.name
        if not name.startswith("xls_") or not name.lower().endswith((".png", ".jpg", ".jpeg")):
            continue
        key = name[4:].rsplit(".", 1)[0]
        mapping[key] = f"{XLS_IMG_DIR}/{name}"
    if "V623040" in mapping:
        mapping.setdefault("V623040A", mapping["V623040"])
        mapping.setdefault("V623040H", mapping["V623040"])
    return mapping


def extract_ats_images():
    out_dir = BASE / ATS_IMG_DIR
    mapping = {}
    if not out_dir.exists():
        return mapping
    # Collect per style, pick the lowest trailing index
    by_style = {}
    for f in out_dir.iterdir():
        if not f.is_file():
            continue
        name = f.name
        if not name.startswith("ATS_") or not name.lower().endswith((".png", ".jpg", ".jpeg")):
            continue
        stem = name[4:].rsplit(".", 1)[0]  # style_color_index
        parts = stem.split("_")
        if len(parts) < 2:
            continue
        style = parts[0]
        idx = 0
        if len(parts) >= 3 and parts[-1].isdigit():
            idx = int(parts[-1])
        by_style.setdefault(style, []).append((idx, name))
    for style, cands in by_style.items():
        cands.sort()
        mapping[style] = f"{ATS_IMG_DIR}/{cands[0][1]}"
    return mapping

image_map = extract_excel_images()
pdf_map = extract_pdf_images()
# Agregar imagenes del nuevo PDF sin sobrescribir las existentes
for _k, _v in extract_pdf_images2().items():
    pdf_map.setdefault(_k, _v)

xls_map = extract_xls_images()
ats_map = extract_ats_images()


def extract_imagenes():
    out_dir = BASE / "imagenes"
    mapping = {}
    if not out_dir.exists():
        return mapping
    for f in out_dir.iterdir():
        if not f.is_file():
            continue
        if f.suffix.lower() not in (".png", ".jpg", ".jpeg"):
            continue
        stem = f.stem
        if stem not in mapping:
            mapping[stem] = f"imagenes/{f.name}"
    return mapping


imagenes_map = extract_imagenes()

df = pd.read_excel(BASE / "Inventario_con_imagenes.xlsx", sheet_name="Data")

# Columna "Color_Name" mapeada desde Inventario-1.xlsx por Style + Color_code
try:
    _inv1 = pd.read_excel(BASE / "sources" / "Inventario-1.xlsx", sheet_name="Data")
    _cn = {}
    for _, _r in _inv1.iterrows():
        _k = (str(_r.get("Style", "")).strip(), str(_r.get("Color_code", "")).strip().upper())
        _v = _r.get("Color Name")
        if _k not in _cn and _v is not None and str(_v).strip() != "" and str(_v).strip().lower() != "nan":
            _cn[_k] = str(_v).strip()
    df["Color_Name"] = df.apply(
        lambda r: _cn.get((str(r.get("Style", "")).strip(),
                           str(r.get("Color_code", "")).strip().upper()), ""),
        axis=1,
    )
except Exception:
    df["Color_Name"] = ""



# Mapeos manuales de imagenes por estilo + color (sobrescriben el resto)
MANUAL_IMAGE_OVERRIDES = {
    ("V322453M", "STO"): "V322453M_COLOR_STONE.png",
}

def combined_image(r):
    _style = str(r.get("Style", "")).strip()
    _color = str(r.get("Color_code", "")).strip().upper()
    if (_style, _color) in MANUAL_IMAGE_OVERRIDES:
        return MANUAL_IMAGE_OVERRIDES[(_style, _color)]
    # Imagen descargada de Rainforest/Amazon para este style
    if _style == "119004414":
        return "imagenes/119004414.png"
    if _style == "117004093":
        return "imagenes/117004093.jpg"
    excel_fname = image_map.get(
        (str(r.get("TLGDIV_NAME", "")).strip(), str(r.get("Style", "")).strip()), ""
    )
    if excel_fname:
        return f"{EXCEL_IMG_DIR}/{excel_fname}"
    pdf_fname = pdf_map.get(str(r.get("Style", "")).strip(), "")
    if pdf_fname:
        return f"{PDF_IMG_DIR}/{pdf_fname}"
    xls_fname = xls_map.get(str(r.get("Style", "")).strip(), "")
    if xls_fname:
        return xls_fname
    ats_fname = ats_map.get(str(r.get("Style", "")).strip(), "")
    if ats_fname:
        return ats_fname
    # Imagenes de la carpeta \imagenes, solo si no hay otra ya mapeada.
    img_fname = imagenes_map.get(str(r.get("Style", "")).strip())
    if img_fname:
        return img_fname
    return ""


# Eliminar filas en blanco / totales / (en blanco)
def _is_excluded_row(r):
    excluded = {"", "(blank)", "nan", "(en blanco)", "total general"}
    tl = str(r.get("TLGDIV_NAME", "")).strip().lower()
    st = str(r.get("Style", "")).strip().lower()
    return tl in excluded or st in excluded

df = df[~df.apply(_is_excluded_row, axis=1)].reset_index(drop=True)

df[IMAGE_COL] = df.apply(combined_image, axis=1)

# Ordenar por proveedor (TLGDIV_NAME) y luego por numero de estilo (Style)
df = df.sort_values(["TLGDIV_NAME", "Style"], kind="mergesort", na_position="last").reset_index(drop=True)


cols = ["#", "TLGDIV_NAME", "Style", "Style_Name", "Color_code", "Color_Name", "Size_Desc", "OTS_INV", IMAGE_COL]


def is_red(row):
    return (str(row.get("TLGDIV_NAME", "")).strip().lower() in ("", "(blank)", "nan")
            or str(row.get("Style", "")).strip().lower() in ("", "(blank)", "nan"))


def unique_vals(col):
    vals = []
    seen = set()
    for v in df[col].tolist():
        if v is None or (isinstance(v, float) and pd.isna(v)):
            s = ""
        elif col == "OTS_INV":
            try:
                f = float(v)
                s = str(int(f)) if f.is_integer() else str(f)
            except (ValueError, TypeError):
                s = str(v).strip()
        elif col == IMAGE_COL:
            s = "con imagen" if (v is not None and str(v).strip() != "" and not (isinstance(v, float) and pd.isna(v))) else "sin imagen"
        else:
            s = str(v).strip()
        if s not in seen:
            seen.add(s)
            vals.append(s)
    return vals


dropdowns = {}
for c in cols:
    if c == "#":
        dropdowns[c] = ""
        continue
    if c == IMAGE_COL:
        vals = ["con imagen", "sin imagen"]
    else:
        vals = unique_vals(c)
    boxes = []
    boxes.append(f'<input type="text" class="filtersearch" placeholder="Buscar...">')
    boxes.append(f'<label class="fitem"><input type="checkbox" class="colchk" data-col="{c}" value="__ALL__" checked> (Todos)</label>')
    for v in vals:
        esc = html.escape(str(v), quote=True)
        boxes.append(f'<label class="fitem"><input type="checkbox" class="colchk" data-col="{c}" value="{esc}" checked> {html.escape(str(v))}</label>')
    dropdowns[c] = "\n".join(boxes)

rows_html = []
for i, (_, r) in enumerate(df.iterrows(), start=1):
    red = is_red(r)
    tr_class = ' class="red"' if red else ""
    cells = []
    data_attrs = []
    for c in cols:
        if c == "#":
            val = i
            s = str(i)
        else:
            val = r.get(c)
            if val is None or (isinstance(val, float) and pd.isna(val)):
                s = ""
            elif c == "OTS_INV":
                try:
                    f = float(val)
                    s = str(int(f)) if f.is_integer() else str(f)
                except (ValueError, TypeError):
                    s = str(val).strip()
            elif c == IMAGE_COL:
                s = "con imagen" if (val is not None and str(val).strip() != "" and not (isinstance(val, float) and pd.isna(val))) else "sin imagen"
            else:
                s = str(val).strip()
        data_attrs.append(f'data-{c}="{html.escape(s, quote=True)}"')
    for c in cols:
        if c == "#":
            cells.append(f"<td>{i}</td>")
            continue
        val = r.get(c)
        if c == IMAGE_COL and isinstance(val, str) and val.strip():
            _v = html.escape(str(val), quote=True)
            img = f'<img src="{_v}" loading="lazy" alt="{html.escape(str(val))}" class="thumb" data-full="{_v}" onclick="openModal(this.dataset.full)" onerror="this.style.display=\'none\'">'
            cells.append(f"<td data-col=\"{c}\">{img}</td>")
        else:
            if val is None or (isinstance(val, float) and pd.isna(val)):
                v = ""
            elif c == "OTS_INV":
                try:
                    f = float(val)
                    v = str(int(f)) if f.is_integer() else str(f)
                except (ValueError, TypeError):
                    v = str(val)
            else:
                v = str(val)
            cells.append(f"<td data-col=\"{c}\">{html.escape(str(v))}</td>")
    rows_html.append(f"<tr{tr_class} {' '.join(data_attrs)}>{''.join(cells)}</tr>")

header_cells = []
for c in cols:
    if c == "#":
        header_cells.append(f"<th data-col='#'>#<div class='col-resizer' data-col='#'></div></th>")
    else:
        header_cells.append(
            f"<th data-col='{c}'>{c} <span class='sortbtn' data-col='{c}' onclick='sortTable(this)' style='cursor:pointer'>↕</span><br><div style='height:10px; line-height:10px;'>&nbsp;</div><button class='filterbtn' data-col='{c}'>Filtrar</button><div class='filterpanel' data-col='{c}' style='display:none'>{dropdowns[c]}</div><div class='col-resizer' data-col='{c}'></div></th>"
        )

html_doc = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<title>Galeria de imagenes - Inventario</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 20px; }}
table {{ border-collapse: collapse; width: 100%; table-layout: fixed; }}
th, td {{ border: 1px solid #ccc; padding: 6px; vertical-align: top; font-size: 12px; }}
th {{ background: #333; color: #fff; position: sticky; top: 0; }}
tr.red {{ background: #ffcccc; }}
th.filtered {{ background: #00bcd4 !important; color: #000 !important; }}
.col-resizer {{ position: absolute; right: 0; top: 0; bottom: 0; width: 6px; cursor: col-resize; z-index: 6; background: transparent; }}
.col-resizer:hover {{ background: rgba(255,255,255,0.35); }}
th[data-col="#"], td[data-col="#"] {{ width: 30px; min-width: 30px; }}
th[data-col="Color_code"], td[data-col="Color_code"] {{ width: 80px; min-width: 80px; }}
th[data-col="Size_Desc"], td[data-col="Size_Desc"] {{ width: 70px; min-width: 70px; }}
th[data-col="OTS_INV"], td[data-col="OTS_INV"] {{ width: 55px; min-width: 55px; }}
th[data-col="TLGDIV_NAME"], td[data-col="TLGDIV_NAME"] {{ width: 90px; min-width: 90px; }}
th[data-col="Style"], td[data-col="Style"] {{ width: 70px; min-width: 70px; }}
th[data-col="Style_Name"], td[data-col="Style_Name"] {{ width: 176px; min-width: 176px; }}
th[data-col="Color_Name"], td[data-col="Color_Name"] {{ width: 80px; min-width: 80px; }}
th[data-col="Imagen"], td[data-col="Imagen"] {{ width: 130px; max-width: 130px; }}
img {{ max-width: 120px; max-height: 120px; object-fit: contain; }}
.sortbtn {{ cursor: pointer; font-size: 12px; color: #fff; user-select: none; }}
th {{ position: relative; }}
.filterbtn {{ font-size: 11px; padding: 2px 6px; cursor: pointer; }}
.filterpanel {{ position: absolute; top: 100%; left: 0; z-index: 1000; background: #fff; color: #000; border: 1px solid #ccc; padding: 6px; max-height: 260px; overflow: auto; text-align: left; min-width: 160px; box-shadow: 0 2px 8px rgba(0,0,0,0.15); }}
.fitem {{ display: block; white-space: nowrap; font-size: 11px; color: #000; }}
.filtersearch {{ width: 95%; margin-bottom: 4px; font-size: 11px; }}
#search {{ margin-bottom: 10px; padding: 8px; width: 300px; }}
select.colfilter {{ width: 100%; font-size: 11px; }}
.thumb {{ cursor: zoom-in; }}
.modal {{ display: none; position: fixed; inset: 0; background: rgba(0,0,0,0.85); z-index: 9999; align-items: center; justify-content: center; flex-direction: column; padding: 20px; }}
.modal.open {{ display: flex; }}
.modal img {{ max-width: 90vw; max-height: 75vh; transition: transform .15s ease; transform-origin: center center; }}
.modal-close {{ position: absolute; top: 12px; right: 20px; color: #fff; font-size: 32px; cursor: pointer; line-height: 1; }}
.modal-controls {{ margin-top: 12px; }}
.modal-controls button {{ margin: 0 4px; padding: 6px 14px; font-size: 13px; cursor: pointer; }}
</style>
</head>
<body>
<h2>Galeria de imagenes - Inventario ({len(df)} filas) <span id="rowCount" style="font-size:13px; font-weight:normal;">(mostrando {len(df)})</span></h2>
<input type="text" id="search" placeholder="Buscar texto en toda la fila...">
<button id="resetFilters" style="margin-left:10px; padding:8px 12px; cursor:pointer;">Reset filters</button>
<table id="tbl">
<thead><tr>{''.join(header_cells)}</tr></thead>
<tbody>
{'\n'.join(rows_html)}
</tbody>
</table>
<div id="modal" class="modal" onclick="if(event.target===this) closeModal()">
  <span class="modal-close" onclick="closeModal()">&times;</span>
  <img id="modalImg" src="" alt="">
  <div class="modal-controls">
    <button onclick="setZoom(1)">1x</button>
    <button onclick="setZoom(2)">2x</button>
    <button onclick="setZoom(3)">3x</button>
  </div>
</div>
<script>
const COLS = ["#", "TLGDIV_NAME", "Style", "Style_Name", "Color_code", "Color_Name", "Size_Desc", "OTS_INV", "Imagen"];
const input=document.getElementById('search');
function rowPassesFilters(tr, filtersToUse, q){{
  if(!tr.textContent.toLowerCase().includes(q)) return false;
  for(const col in filtersToUse){{
    const val = tr.getAttribute('data-'+col);
    const f = filtersToUse[col];
    if (f instanceof Set) {{
      if (f.size === 0 || !f.has(val)) return false;
    }} else if (f && val !== f) return false;
  }}
  return true;
}}

function applyFilters(){{
  const q=input.value.toLowerCase();
  const filters={{}};
  const filteredCols = new Set();
  document.querySelectorAll('.filterpanel').forEach(panel=>{{
    const col = panel.dataset.col;
    const boxes = panel.querySelectorAll('.colchk');
    const checked = [...boxes].filter(b=>b.checked && b.value !== '__ALL__').map(b=>b.value);
    const allBox = panel.querySelector('.colchk[value="__ALL__"]');
    if (allBox && allBox.checked) {{
      // todos seleccionados -> sin filtro
    }} else {{
      filteredCols.add(col);
      if (checked.length === 0) {{
        filters[col] = new Set(); // nada seleccionado -> ocultar todo
      }} else {{
        filters[col] = new Set(checked);
      }}
    }}
  }});

  // Filtro cruzado (faceted)
  const allRows = Array.from(document.querySelectorAll('#tbl tbody tr'));
  document.querySelectorAll('.filterpanel').forEach(panel=>{{
    const targetCol = panel.dataset.col;
    const otherFilters = {{}};
    for(const c in filters){{
      if (c !== targetCol) otherFilters[c] = filters[c];
    }}
    const available = new Set();
    allRows.forEach(tr=>{{
      if (rowPassesFilters(tr, otherFilters, q)) {{
        available.add(tr.getAttribute('data-'+targetCol));
      }}
    }});
    panel.querySelectorAll('.colchk').forEach(cb=>{{
      if (cb.value === '__ALL__') return;
      const isAvailable = available.has(cb.value);
      cb.disabled = !isAvailable;
      const lbl = cb.parentElement;
      if (lbl) {{
        lbl.style.opacity = isAvailable ? '1' : '0.45';
        lbl.style.pointerEvents = isAvailable ? '' : 'none';
      }}
    }});
  }});

  // resaltar columnas filtradas en cyan
  document.querySelectorAll('th').forEach(th=>{{
    const btn = th.querySelector('.filterbtn');
    const col = btn ? btn.dataset.col : null;
    if (col && filteredCols.has(col)) th.classList.add('filtered');
    else th.classList.remove('filtered');
  }});
  document.querySelectorAll('#tbl tbody tr').forEach(tr=>{{
    let show = tr.textContent.toLowerCase().includes(q);
    for(const col in filters){{
      const val = tr.getAttribute('data-'+col);
      if (filters[col] instanceof Set) {{
        if (filters[col].size === 0 || !filters[col].has(val)) {{ show=false; break; }}
      }} else if (filters[col] && val !== filters[col]) {{ show=false; break; }}
    }}
    tr.style.display = show ? '' : 'none';
  }});
  const visible = Array.from(document.querySelectorAll('#tbl tbody tr')).filter(tr=>tr.style.display!=='none').length;
  const rc = document.getElementById('rowCount');
  if(rc) rc.textContent = '(mostrando ' + visible + ' de {len(df)})';
}}

// === Resize de columnas con el mouse ===
let _resizing = null;
document.addEventListener('mousedown', function(e){{
  const handle = (e.target && e.target.classList && e.target.classList.contains('col-resizer')) ? e.target : null;
  if (!handle) return;
  const th = handle.parentElement;
  _resizing = {{ th: th, startX: e.clientX, startWidth: th.getBoundingClientRect().width }};
  document.body.style.cursor = 'col-resize';
  document.body.style.userSelect = 'none';
  e.preventDefault();
  e.stopPropagation();
}});
document.addEventListener('mousemove', function(e){{
  if (!_resizing) return;
  const delta = e.clientX - _resizing.startX;
  const newWidth = Math.max(60, _resizing.startWidth + delta);
  _resizing.th.style.width = newWidth + 'px';
  _resizing.th.style.minWidth = newWidth + 'px';
}});
document.addEventListener('mouseup', function(){{
  if (!_resizing) return;
  const widths = {{}};
  document.querySelectorAll('th[data-col]').forEach(function(th){{ widths[th.dataset.col] = th.getBoundingClientRect().width; }});
  try {{ localStorage.setItem('colWidths', JSON.stringify(widths)); }} catch(err) {{}}
  _resizing = null;
  document.body.style.cursor = '';
  document.body.style.userSelect = '';
}});
try {{
  const saved = JSON.parse(localStorage.getItem('colWidths') || '{{}}');
  Object.keys(saved).forEach(function(col){{
    const th = document.querySelector('th[data-col="'+col+'"]');
    if (th) {{ th.style.width = saved[col] + 'px'; th.style.minWidth = saved[col] + 'px'; }}
  }});
}} catch(err) {{}}

input.addEventListener('input', applyFilters);
// inicializar disponibilidad cruzada al cargar
applyFilters();
const resetBtn = document.getElementById('resetFilters');
if (resetBtn) {{
  resetBtn.addEventListener('click', function(){{
    input.value = '';
    document.querySelectorAll('.filterpanel .colchk').forEach(function(b){{ b.checked = true; }});
    applyFilters();
  }});
}}

let sortState = {{ col: null, asc: true }};
function sortTable(btn){{
  const col = btn.dataset.col;
  const asc = (sortState.col === col) ? !sortState.asc : true;
  sortState = {{ col: col, asc: asc }};
  const tbody = document.querySelector('#tbl tbody');
  const rows = Array.from(tbody.querySelectorAll('tr'));
  const idx = COLS.indexOf(col);
  if (idx < 0) return;
  rows.sort(function(a,b){{
    const va = (a.cells[idx] ? a.cells[idx].textContent : '').trim();
    const vb = (b.cells[idx] ? b.cells[idx].textContent : '').trim();
    const na = parseFloat(va.replace(/,/g,''));
    const nb = parseFloat(vb.replace(/,/g,''));
    if (!isNaN(na) && !isNaN(nb)) return asc ? (na - nb) : (nb - na);
    return asc ? va.localeCompare(vb) : vb.localeCompare(va);
  }});
  rows.forEach(function(r){{ tbody.appendChild(r); }});
  document.querySelectorAll('.sortbtn').forEach(function(b){{ b.textContent = '\u2195'; }});
  btn.textContent = asc ? '\u25B2' : '\u25BC';
  applyFilters();
}}


// Toggle de paneles de filtro
document.addEventListener('click', function(e){{
  if (e.target && e.target.classList && e.target.classList.contains('filterbtn')){{
    const col = e.target.dataset.col;
    const panel = document.querySelector('.filterpanel[data-col="'+col+'"]');
    if (panel) panel.style.display = panel.style.display === 'none' ? 'block' : 'none';
  }} else if (!e.target.closest('.filterpanel') && !(e.target.classList && e.target.classList.contains('filterbtn'))){{
    document.querySelectorAll('.filterpanel').forEach(function(p){{ p.style.display='none'; }});
  }}
}});
// Select-all y cambio por checkbox
document.addEventListener('change', function(e){{
  if (e.target && e.target.classList && e.target.classList.contains('colchk') && e.target.value === '__ALL__'){{
    const col = e.target.dataset.col;
    const panel = document.querySelector('.filterpanel[data-col="'+col+'"]');
    if (panel) panel.querySelectorAll('.colchk').forEach(function(b){{ b.checked = e.target.checked; }});
    applyFilters();
  }} else if (e.target && e.target.classList && e.target.classList.contains('colchk')){{
    const col = e.target.dataset.col;
    const panel = document.querySelector('.filterpanel[data-col="'+col+'"]');
    if (panel){{
      const allBox = panel.querySelector('.colchk[value="__ALL__"]');
      const boxes = [...panel.querySelectorAll('.colchk')].filter(function(b){{ return b.value !== '__ALL__'; }});
      if (allBox) allBox.checked = boxes.every(function(b){{ return b.checked; }});
    }}
    applyFilters();
  }}
}});
// Busqueda dentro del panel
document.addEventListener('input', function(e){{
  if (e.target && e.target.classList && e.target.classList.contains('filtersearch')){{
    const panel = e.target.closest('.filterpanel');
    const q = e.target.value.toLowerCase();
    panel.querySelectorAll('.fitem').forEach(function(lbl){{
      const txt = lbl.textContent.toLowerCase();
      lbl.style.display = txt.includes(q) ? '' : 'none';
    }});
  }}
}});

// Modal de imagen con zoom
let zoomLevel = 1;
function openModal(src){{
  const modal = document.getElementById('modal');
  const img = document.getElementById('modalImg');
  if(!modal || !img) return;
  img.src = src;
  zoomLevel = 1;
  img.style.transform = 'scale(1)';
  modal.classList.add('open');
}}
function closeModal(){{
  const modal = document.getElementById('modal');
  const img = document.getElementById('modalImg');
  if(modal) modal.classList.remove('open');
  if(img){{ img.src = ''; img.style.transform = 'scale(1)'; }}
  zoomLevel = 1;
}}
function setZoom(z){{
  zoomLevel = z;
  const img = document.getElementById('modalImg');
  if(img) img.style.transform = 'scale(' + z + ')';
}}
document.addEventListener('keydown', function(e){{ if(e.key === 'Escape') closeModal(); }});
const _modalEl = document.getElementById('modal');
if(_modalEl){{
  _modalEl.addEventListener('wheel', function(e){{
    if(!_modalEl.classList.contains('open')) return;
    e.preventDefault();
    zoomLevel = Math.min(5, Math.max(1, zoomLevel + (e.deltaY < 0 ? 0.5 : -0.5)));
    setZoom(zoomLevel);
  }}, {{ passive: false }});
}}


</script>
</body>
</html>
"""
out = BASE / "index.html"
out.write_text(html_doc, encoding="utf-8")
print("Generado:", out)
