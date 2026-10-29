from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from image_sources import build_maps, combined_image

BASE = Path(__file__).resolve().parent
EXCEL_IN = BASE / "sources" / "Inventario.xlsx"
EXCEL_IN1 = BASE / "sources" / "Inventario-1.xlsx"
EXCEL_OUT = BASE / "Inventario_con_imagenes.xlsx"
SHEET_NAME = "Data"
IMAGE_COL = "Imagen"

SOURCE_COLUMNS = ["BRAND_NAME", "Style", "Style_Name", "Color_code",
                  "Dm_Pk", "Lbl_code", "Size_Desc", "QTY"]
REQUIRED_COLS = ["BRAND_NAME", "Style", "Style_Name", "Color_code"]


def is_missing(v):
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    s = str(v).strip()
    return s == "" or s.lower() in ("nan", "(blank)", "(en blanco)", "none", "nat")


def load_inventory():
    df = pd.read_excel(EXCEL_IN, sheet_name=SHEET_NAME, header=2, usecols="A:H")
    if df.shape[1] != len(SOURCE_COLUMNS):
        raise SystemExit(f"{EXCEL_IN.name}: se esperaban {len(SOURCE_COLUMNS)} columnas, hay {df.shape[1]}")
    df.columns = SOURCE_COLUMNS

    missing_mask = df[REQUIRED_COLS].apply(lambda c: c.map(is_missing)).any(axis=1)

    maps = build_maps()
    df[IMAGE_COL] = df.apply(lambda r: combined_image(r, maps), axis=1)

    try:
        inv1 = pd.read_excel(EXCEL_IN1, sheet_name=SHEET_NAME)
        cn = {}
        for _, r in inv1.iterrows():
            k = (str(r.get("Style", "")).strip(), str(r.get("Color_code", "")).strip().upper())
            v = r.get("Color Name")
            if k not in cn and v is not None and str(v).strip() != "" and str(v).strip().lower() != "nan":
                cn[k] = str(v).strip()
        df["Color_Name"] = df.apply(
            lambda r: cn.get((str(r.get("Style", "")).strip(),
                              str(r.get("Color_code", "")).strip().upper()), ""),
            axis=1,
        )
    except Exception:
        df["Color_Name"] = ""

    return df, missing_mask


def write_excel(df, missing_mask):
    out = df[SOURCE_COLUMNS + [IMAGE_COL]].copy()
    out.to_excel(EXCEL_OUT, sheet_name=SHEET_NAME, index=False, engine="openpyxl")

    wb = load_workbook(EXCEL_OUT)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    col_index = {name: i + 1 for i, name in enumerate(header)}

    red_font = Font(color="FF0000", bold=True)
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    for idx in df.index[missing_mask].tolist():
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=idx + 2, column=c)
            cell.font = red_font
            cell.fill = red_fill

    cidx = col_index[IMAGE_COL]
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=cidx)
        if cell.value:
            cell.hyperlink = str(cell.value)

    widths = {"Style_Name": 34, "Color_Name": 21, "BRAND_NAME": 21, "Style": 21, IMAGE_COL: 24}
    for c in range(1, ws.max_column + 1):
        name = header[c - 1] if c - 1 < len(header) else ""
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 22)

    wb.save(EXCEL_OUT)
    print(f"Excel escrito: {EXCEL_OUT}")
