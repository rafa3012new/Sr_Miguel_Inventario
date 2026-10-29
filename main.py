"""
Image Generator API - builds Inventario_con_imagenes.xlsx with a single
'Imagen' column pointing to the same images used by the HTML gallery
(imagenes_excel / imagenes_pdf / imagenes_xls), instead of downloading
images from the internet.

Usage:
    python main.py
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from image_sources import build_maps, combined_image

BASE_DIR = Path(__file__).resolve().parent
EXCEL_IN = BASE_DIR / "sources" / "Inventario.xlsx"
EXCEL_OUT = BASE_DIR / "Inventario_con_imagenes.xlsx"
SHEET_NAME = "Data"

IMAGE_COL = "Imagen"

REQUIRED_COLS = ["TLGDIV_NAME", "Style", "Style_Name", "Color_code"]


def is_missing(v):
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    s = str(v).strip()
    return s == "" or s.lower() in ("nan", "(blank)", "(en blanco)", "none", "nat")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="Compatibilidad (ya no descarga).")
    parser.add_argument("--limit", type=int, default=None, help="Compatibilidad.")
    args = parser.parse_args()

    if not EXCEL_IN.exists():
        raise SystemExit(f"No se encontro {EXCEL_IN}")

    # Read Excel
    df = pd.read_excel(EXCEL_IN, sheet_name=SHEET_NAME, header=2, usecols="A:H")
    df.columns = ["TLGDIV_NAME", "Style", "Style_Name", "Color_code",
                  "Dm_Pk", "Lbl_code", "Size_Desc", "OTS_INV"]

    missing_mask = df[REQUIRED_COLS].apply(lambda c: c.map(is_missing)).any(axis=1)
    missing_idx = df.index[missing_mask].tolist()

    print(f"Filas totales: {len(df)}")
    print(f"Filas con datos faltantes (rojo): {len(missing_idx)}")

    maps = build_maps()
    df[IMAGE_COL] = df.apply(lambda r: combined_image(r, maps), axis=1)

    out_df = df.copy()

    out_df.to_excel(EXCEL_OUT, sheet_name=SHEET_NAME, index=False, engine="openpyxl")
    print(f"Excel escrito: {EXCEL_OUT}")

    # Post-process: red rows + hyperlinks for image column
    wb = load_workbook(EXCEL_OUT)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    col_index = {name: i + 1 for i, name in enumerate(header)}

    red_font = Font(color="FF0000", bold=True)
    red_fill = PatternFill("solid", fgColor="FFC7CE")

    for idx in missing_idx:
        row_num = idx + 2
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_num, column=c)
            cell.font = red_font
            cell.fill = red_fill

    if IMAGE_COL in col_index:
        cidx = col_index[IMAGE_COL]
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=cidx)
            val = cell.value
            if val:
                rel = str(val)
                cell.hyperlink = rel
                cell.value = val

    for c in range(1, ws.max_column + 1):
        _cname = header[c - 1] if c - 1 < len(header) else ""
        if _cname == "Style_Name":
            ws.column_dimensions[get_column_letter(c)].width = 34
        elif _cname == "Color_Name":
            ws.column_dimensions[get_column_letter(c)].width = 21
        elif _cname == "TLGDIV_NAME":
            ws.column_dimensions[get_column_letter(c)].width = 21
        elif _cname == "Style":
            ws.column_dimensions[get_column_letter(c)].width = 21
        elif _cname == "Imagen":
            ws.column_dimensions[get_column_letter(c)].width = 24
        else:
            ws.column_dimensions[get_column_letter(c)].width = 22

    wb.save(EXCEL_OUT)
    print("Post-procesado listo (filas rojas + hyperlinks).")


if __name__ == "__main__":
    main()